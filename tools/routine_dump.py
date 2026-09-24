#!/usr/bin/env python3
"""ルーティンチェックシートの中身を、検証用にJSONへ書き出す。

本物のスプレッドシートは読めないので、Excelに書き出したものを使う。
シートの形（1行目に開催日／項目名はB〜E列）が変わっていないかを、
実物のデータで確かめるための下ごしらえ。

    python3 tools/routine_dump.py <xlsx> <出力.json>
"""
import datetime
import json
import sys

import openpyxl

ROWS = 90


def cell(v):
    if v is None:
        return ''
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y/%m/%d')
    return v


def main():
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    out = {}
    for name in wb.sheetnames:
        if 'ルーティンチェックシート' not in name or not name.startswith('【'):
            continue
        ws = wb[name]
        cols = ws.max_column
        grid = []
        for r in range(1, min(ROWS, ws.max_row) + 1):
            grid.append([cell(ws.cell(r, c).value) for c in range(1, cols + 1)])
        out[name] = grid
    json.dump(out, open(sys.argv[2], 'w'), ensure_ascii=False)
    print('書き出し: %d シート' % len(out))
    for k, v in out.items():
        print('  %-28s %d行 x %d列' % (k, len(v), len(v[0]) if v else 0))


if __name__ == '__main__':
    main()
