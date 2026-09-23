# -*- coding: utf-8 -*-
"""BNI ビジタースライドのテンプレートpptxを作る。
差し込み先は p:cNvPr/@id で特定されるので、シェイプIDを所定の値に固定する。
  ビジタープレゼン(1人1枚) : 25=氏名+様 / 27=会社名 / 29=【カテゴリー】
  ビジター紹介・代理紹介(3人1枚):
      1人目 19=カテゴリー 20=招待者 21=氏名+様
      2人目 25=カテゴリー 26=招待者 28=氏名+様   ※2人目の氏名は27ではなく28
      3人目 31=カテゴリー 32=招待者 33=氏名+様
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

RED   = RGBColor(0xC0, 0x00, 0x00)
NAVY  = RGBColor(0x16, 0x23, 0x3F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY  = RGBColor(0x44, 0x44, 0x44)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
FONT  = 'Meiryo'

W, H = Inches(13.333), Inches(7.5)


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs, prs.slides.add_slide(prs.slide_layouts[6])   # 白紙


def box(slide, l, t, w, h, text, size, *, color=NAVY, bold=False,
        align=PP_ALIGN.CENTER, shape_id=None, runs=None, name=None):
    """テキストボックスを1つ置く。runs を渡すと複数ランで作る（【】の分割用）。"""
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    for part in (runs if runs else [text]):
        r = p.add_run()
        r.text = part
        f = r.font
        f.size, f.bold, f.name = Pt(size), bold, FONT
        f.color.rgb = color
    if shape_id is not None:
        tb._element.nvSpPr.cNvPr.set('id', str(shape_id))
    if name:
        tb._element.nvSpPr.cNvPr.set('name', name)
    return tb


def band(slide, l, t, w, h, fill, shape_id=None):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    if shape_id is not None:
        sh._element.nvSpPr.cNvPr.set('id', str(shape_id))
    return sh


def header(slide, title):
    band(slide, 0, 0, W, Inches(1.05), RED, shape_id=2)
    box(slide, Inches(0.5), Inches(0.1), W - Inches(1.0), Inches(0.85), title, 26,
        color=WHITE, bold=True, shape_id=3, name='見出し（自由に変更可）')


# ---------- ビジタープレゼン（1人1枚） ----------
def make_presen(path):
    prs, s = new_deck()
    header(s, 'ビジタープレゼンテーション')
    # 【カテゴリー】は 3 ラン構成。差し込み時に「【」「値」「】」へ個別に入る
    box(s, Inches(1.0), Inches(1.6), W - Inches(2.0), Inches(1.0),
        None, 28, color=RED, bold=True, shape_id=29,
        runs=['【', 'カテゴリー', '】'], name='カテゴリー')
    box(s, Inches(1.0), Inches(2.8), W - Inches(2.0), Inches(1.8),
        'お名前 様', 60, bold=True, shape_id=25, name='氏名')
    box(s, Inches(1.0), Inches(4.8), W - Inches(2.0), Inches(1.0),
        '会社名', 28, color=GRAY, shape_id=27, name='会社名')
    band(s, 0, H - Inches(0.35), W, Inches(0.35), NAVY, shape_id=4)
    prs.save(path)


# ---------- ビジター紹介／代理紹介（3人1枚） ----------
def make_group(path, title, ids):
    prs, s = new_deck()
    header(s, title)
    colw = (W - Inches(1.6)) / 3
    gap = Inches(0.2)
    for i, (cid, iid, nid) in enumerate(ids):
        left = Inches(0.6) + (colw + gap) * i - gap * i
        left = Inches(0.6) + i * (colw)
        band(s, left + Inches(0.08), Inches(1.5), colw - Inches(0.16), Inches(5.2), LIGHT, shape_id=10 + i)
        box(s, left + Inches(0.18), Inches(1.7), colw - Inches(0.36), Inches(0.9),
            'カテゴリー', 16, color=RED, bold=True, shape_id=cid, name='カテゴリー%d' % (i + 1))
        box(s, left + Inches(0.18), Inches(2.6), colw - Inches(0.36), Inches(1.2),
            'お名前 様', 30, bold=True, shape_id=nid, name='氏名%d' % (i + 1))
        # 招待者。「ご招待」の固定ラベルは置かない。空き枠のときにラベルだけ残ってしまうため
        box(s, left + Inches(0.18), Inches(4.0), colw - Inches(0.36), Inches(0.8),
            '招待者', 14, color=GRAY, shape_id=iid, name='招待者%d' % (i + 1))
    band(s, 0, H - Inches(0.35), W, Inches(0.35), NAVY, shape_id=5)
    prs.save(path)


GROUP_IDS = [(19, 20, 21), (25, 26, 28), (31, 32, 33)]
import os
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'templates')
os.makedirs(OUT, exist_ok=True)
make_presen(os.path.join(OUT, 'BNI_テンプレート_ビジタープレゼン.pptx'))
make_group(os.path.join(OUT, 'BNI_テンプレート_ビジター紹介.pptx'), 'ビジターのご紹介', GROUP_IDS)
make_group(os.path.join(OUT, 'BNI_テンプレート_代理紹介.pptx'), '代理出席者のご紹介', GROUP_IDS)
print('templates/ に3つ作成しました')
